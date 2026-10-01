"""
Mini ERP: exports ERP data from the database into an Excel report.

Usage:
    python export_to_excel.py              # reads MySQL (set DB_* env vars below)
    python export_to_excel.py --demo       # builds a throw-away SQLite copy from the .sql files (no MySQL needed)

Requires: pip install openpyxl mysql-connector-python
"""
import os, re, sys, sqlite3
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
OUT = HERE / "Mini_ERP_Report.xlsx"

SALES_SQL = """SELECT order_id, order_date, customer_name, product_name, category, quantity, unit_price
               FROM v_sales_detail ORDER BY order_date, order_id"""
INVENTORY_SQL = """SELECT p.product_id, p.sku, p.product_name, s.supplier_name, p.stock_qty, p.reorder_level, p.unit_cost
                   FROM products p JOIN suppliers s ON s.supplier_id = p.supplier_id ORDER BY p.product_id"""


def connect(demo):
    if demo:  # build SQLite copy from the same .sql files (strip MySQL-only syntax)
        con = sqlite3.connect(":memory:")
        for name in ("schema.sql", "sample_data.sql"):
            sql = (HERE / name).read_text()
            sql = re.sub(r"(?im)^\s*(DROP DATABASE|CREATE DATABASE|USE)\b.*$", "", sql)
            sql = re.sub(r"(?i)\s*AUTO_INCREMENT", "", sql)
            sql = re.sub(r"(?i)\)\s*ENGINE=InnoDB", ")", sql)
            con.executescript(sql)
        return con
    import mysql.connector
    return mysql.connector.connect(
        host=os.getenv("DB_HOST", "localhost"), user=os.getenv("DB_USER", "root"),
        password=os.getenv("DB_PASSWORD", ""), database=os.getenv("DB_NAME", "mini_erp"))


def fetch(con, sql):
    cur = con.cursor()
    cur.execute(sql)
    rows = [tuple(float(v) if v.__class__.__name__ == "Decimal" else (str(v) if hasattr(v, "isoformat") else v)
                  for v in r) for r in cur.fetchall()]
    cur.close()
    return rows


HDR_FILL = PatternFill("solid", start_color="1F3864")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BASE = Font(name="Arial", size=10)


def header(ws, cols, row=1):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def finish(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.font == Font():  # untouched default font
                c.font = BASE
    ws.freeze_panes = "A2"


def build(sales, inventory):
    wb = Workbook()
    # ---------- Sales_Detail ----------
    sd = wb.active; sd.title = "Sales_Detail"
    header(sd, ["Order ID", "Order Date", "Month", "Customer", "Product", "Category", "Qty", "Unit Price", "Line Total"])
    for r, (oid, dt, cust, prod, cat, qty, price) in enumerate(sales, 2):
        sd.append([oid, dt, dt[:7], cust, prod, cat, qty, price, f"=G{r}*H{r}"])
    n_s = len(sales) + 1
    for r in range(2, n_s + 1):
        sd[f"H{r}"].number_format = sd[f"I{r}"].number_format = '#,##0.00'
    finish(sd, [10, 12, 10, 24, 26, 14, 8, 12, 14])

    # ---------- Inventory ----------
    inv = wb.create_sheet("Inventory")
    header(inv, ["Product ID", "SKU", "Product", "Supplier", "Stock Qty", "Reorder Level", "Unit Cost", "Stock Value", "Status"])
    for r, row in enumerate(inventory, 2):
        inv.append(list(row) + [f"=E{r}*G{r}", f'=IF(E{r}<=F{r},"REORDER","OK")'])
    n_i = len(inventory) + 1
    for r in range(2, n_i + 1):
        inv[f"G{r}"].number_format = inv[f"H{r}"].number_format = '#,##0.00'
    finish(inv, [11, 10, 26, 24, 11, 14, 12, 14, 11])
    from openpyxl.formatting.rule import CellIsRule
    inv.conditional_formatting.add(f"I2:I{n_i}", CellIsRule(operator="equal", formula=['"REORDER"'],
                                   fill=PatternFill("solid", start_color="F8CBAD", end_color="F8CBAD")))

    # ---------- Sales_By_Product (SUMIFS summary, sorted by revenue) ----------
    rev = {}
    for _, _, _, prod, _, qty, price in sales:
        rev[prod] = rev.get(prod, 0) + qty * price
    products = sorted(rev, key=rev.get, reverse=True)
    sp = wb.create_sheet("Sales_By_Product")
    header(sp, ["Product", "Units Sold", "Revenue"])
    for r, p in enumerate(products, 2):
        sp.append([p,
                   f"=SUMIFS(Sales_Detail!$G$2:$G${n_s},Sales_Detail!$E$2:$E${n_s},A{r})",
                   f"=SUMIFS(Sales_Detail!$I$2:$I${n_s},Sales_Detail!$E$2:$E${n_s},A{r})"])
        sp[f"C{r}"].number_format = '#,##0'
    n_p = len(products) + 1
    finish(sp, [28, 12, 14])
    ch = BarChart(); ch.type = "bar"; ch.title = "Revenue by Product"; ch.style = 10
    ch.add_data(Reference(sp, min_col=3, min_row=1, max_row=min(n_p, 11)), titles_from_data=True)
    ch.set_categories(Reference(sp, min_col=1, min_row=2, max_row=min(n_p, 11)))
    ch.legend = None; ch.y_axis.title = "Revenue"; ch.height, ch.width = 9, 18
    ch.x_axis.scaling.orientation = "maxMin"
    ch.x_axis.delete = False; ch.y_axis.delete = False
    sp.add_chart(ch, "E2")

    # ---------- Monthly_Sales ----------
    months = sorted({dt[:7] for _, dt, *_ in sales})
    ms = wb.create_sheet("Monthly_Sales")
    header(ms, ["Month", "Revenue"])
    for r, m in enumerate(months, 2):
        ms.append([m, f"=SUMIFS(Sales_Detail!$I$2:$I${n_s},Sales_Detail!$C$2:$C${n_s},A{r})"])
        ms[f"B{r}"].number_format = '#,##0'
    n_m = len(months) + 1
    finish(ms, [12, 14])
    lc = LineChart(); lc.title = "Monthly Revenue"; lc.style = 12
    lc.add_data(Reference(ms, min_col=2, min_row=1, max_row=n_m), titles_from_data=True)
    lc.set_categories(Reference(ms, min_col=1, min_row=2, max_row=n_m))
    lc.legend = None; lc.height, lc.width = 9, 18
    lc.series[0].smooth = False
    lc.x_axis.delete = False; lc.y_axis.delete = False
    lc.y_axis.title = "Revenue"
    ms.add_chart(lc, "D2")

    # ---------- Summary ----------
    sm = wb.create_sheet("Summary", 0)
    sm["A1"] = "Mini ERP: Sales & Inventory Summary"; sm["A1"].font = Font(name="Arial", bold=True, size=14, color="1F3864")
    kpis = [
        ("Total Revenue", f"=SUM(Sales_Detail!I2:I{n_s})", '#,##0'),
        ("Units Sold", f"=SUM(Sales_Detail!G2:G{n_s})", '#,##0'),
        ("Order Lines", f"=COUNT(Sales_Detail!A2:A{n_s})", '#,##0'),
        ("Inventory Value (at cost)", f"=SUM(Inventory!H2:H{n_i})", '#,##0'),
        ("Products to Reorder", f'=COUNTIF(Inventory!I2:I{n_i},"REORDER")', '0'),
        ("Top Product by Revenue", f"=INDEX(Sales_By_Product!A2:A{n_p},MATCH(MAX(Sales_By_Product!C2:C{n_p}),Sales_By_Product!C2:C{n_p},0))", '@'),
        ("Best Month by Revenue", f"=INDEX(Monthly_Sales!A2:A{n_m},MATCH(MAX(Monthly_Sales!B2:B{n_m}),Monthly_Sales!B2:B{n_m},0))", '@'),
    ]
    for i, (k, f, fmt) in enumerate(kpis, 3):
        sm[f"A{i}"], sm[f"B{i}"] = k, f
        sm[f"A{i}"].font = Font(name="Arial", bold=True); sm[f"B{i}"].font = BASE
        sm[f"B{i}"].number_format = fmt; sm[f"B{i}"].alignment = Alignment(horizontal="right")
    sm["A11"] = "Note: all figures are formulas over the Sales_Detail and Inventory sheets (data exported from MySQL). Amounts in INR."
    sm["A11"].font = Font(name="Arial", italic=True, size=9)
    sm.column_dimensions["A"].width = 30; sm.column_dimensions["B"].width = 26
    wb.save(OUT)
    print("Saved", OUT)


if __name__ == "__main__":
    con = connect("--demo" in sys.argv)
    build(fetch(con, SALES_SQL), fetch(con, INVENTORY_SQL))
